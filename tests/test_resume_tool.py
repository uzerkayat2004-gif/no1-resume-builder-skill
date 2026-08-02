from __future__ import annotations

import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = ROOT / "skills" / "resume-builder"
MODULE_PATH = SKILL_ROOT / "scripts" / "resume_tool.py"
SPEC = importlib.util.spec_from_file_location("resume_tool", MODULE_PATH)
assert SPEC and SPEC.loader
resume_tool = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(resume_tool)


class ResumeToolTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.record = json.loads(
            (SKILL_ROOT / "examples" / "sample-career-record.json").read_text(
                encoding="utf-8"
            )
        )

    def test_sample_record_is_valid_and_gap_is_visible(self) -> None:
        report = resume_tool.validate_record(self.record)
        self.assertTrue(report["valid"], report["errors"])
        self.assertTrue(
            any("Marketplace pricing analytics" in item for item in report["warnings"])
        )

    def test_unverified_claim_blocks_final_render(self) -> None:
        record = copy.deepcopy(self.record)
        record["evidence"][0]["sourceStatus"] = "unverified"
        report = resume_tool.validate_record(record)
        self.assertFalse(report["valid"])
        self.assertTrue(
            any("uses unverified evidence ev-years" in item for item in report["errors"])
        )

    def test_numeric_detail_must_exist_in_evidence(self) -> None:
        record = copy.deepcopy(self.record)
        record["work"][0]["highlights"][0][
            "text"
        ] = "Built a weekly funnel that increased activation by 48%."
        report = resume_tool.validate_record(record)
        self.assertFalse(report["valid"])
        self.assertTrue(
            any("48%" in item and "numeric detail" in item for item in report["errors"])
        )

    def test_render_and_round_trip(self) -> None:
        try:
            import docx  # noqa: F401
            import pypdf  # noqa: F401
            import reportlab  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"Optional renderer dependency unavailable: {exc}")

        with tempfile.TemporaryDirectory() as temp:
            outdir = Path(temp)
            manifest = resume_tool.render_all(
                self.record, outdir, ["docx", "pdf", "txt", "md"]
            )
            self.assertEqual(set(manifest["formats"]), {"docx", "pdf", "txt", "md"})
            for path in manifest["formats"].values():
                self.assertGreater(Path(path).stat().st_size, 0)

            docx_path = Path(manifest["formats"]["docx"])
            pdf_path = Path(manifest["formats"]["pdf"])
            self.assertTrue(
                resume_tool.verify_document(self.record, docx_path)["pass"]
            )
            self.assertTrue(
                resume_tool.verify_document(self.record, pdf_path)["pass"]
            )

            from docx import Document
            from docx.oxml.ns import qn

            document = Document(docx_path)
            self.assertEqual(len(document.tables), 0)
            self.assertTrue(
                all(not paragraph.text for paragraph in document.sections[0].header.paragraphs)
            )
            self.assertTrue(
                all(not paragraph.text for paragraph in document.sections[0].footer.paragraphs)
            )
            bullet_paragraphs = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.style.name == "Resume Bullet"
            ]
            self.assertGreater(len(bullet_paragraphs), 0)
            self.assertTrue(
                all(
                    paragraph._p.pPr is not None
                    and paragraph._p.pPr.find(qn("w:numPr")) is not None
                    for paragraph in bullet_paragraphs
                )
            )

    def test_interview_scaffold_keeps_gap_honest(self) -> None:
        output = resume_tool.interview_markdown(self.record)
        self.assertIn("Marketplace pricing analytics", output)
        self.assertIn("unsupported", output)
        self.assertIn("ev-experiment", output)
        self.assertIn("Do not memorize a fictional script", output)

    def test_domain_router_covers_distinct_occupations(self) -> None:
        cases = {
            "Senior backend software engineer working on cloud systems": "technology-software",
            "Registered nurse in an emergency department": "clinical-healthcare",
            "Hotel front desk guest service supervisor": "retail-hospitality-service",
            "Electrician and field service technician": "construction-trades-field-service",
            "Financial analyst supporting quarterly forecasts": "finance-accounting",
            "High school mathematics teacher": "education-training",
        }
        for query, expected in cases.items():
            with self.subTest(query=query):
                report = resume_tool.route_domain(query)
                self.assertEqual(report["selected"]["id"], expected)

    def test_unknown_domain_uses_general_profile(self) -> None:
        report = resume_tool.route_domain("specialist in an emerging unnamed craft")
        self.assertEqual(report["selected"]["id"], "general")
        self.assertTrue(report["needsConfirmation"])

    def test_scratch_session_asks_no_more_than_three_questions(self) -> None:
        session = resume_tool.new_intake_session(
            journey="scratch",
            role="Registered Nurse",
            region="UAE",
        )
        self.assertEqual(session["target"]["domain"]["id"], "clinical-healthcare")
        self.assertLessEqual(len(session["nextQuestions"]), 3)
        report = resume_tool.intake_status_report(session)
        self.assertEqual(report["progress"]["source"], "not-applicable")
        self.assertEqual(report["progress"]["target"], "complete")

    def test_existing_resume_session_extracts_without_overwriting(self) -> None:
        try:
            import docx  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"Optional renderer dependency unavailable: {exc}")

        with tempfile.TemporaryDirectory() as temp:
            outdir = Path(temp)
            manifest = resume_tool.render_all(self.record, outdir, ["docx"])
            source = Path(manifest["formats"]["docx"])
            source_hash = resume_tool._file_digest(source)
            session = resume_tool.new_intake_session(
                journey="existing-resume",
                source=source,
                role="Product Analyst",
                region="India",
            )
            self.assertEqual(resume_tool._file_digest(source), source_hash)
            self.assertGreater(session["sources"][0]["extractedCharacters"], 100)
            self.assertFalse(session["answers"]["sourceReviewed"])
            self.assertTrue(
                any(
                    item["id"] == "source.confirm"
                    for item in session["nextQuestions"]
                )
            )

    def test_skill_is_progressively_disclosed(self) -> None:
        skill_lines = (SKILL_ROOT / "SKILL.md").read_text(encoding="utf-8").splitlines()
        self.assertLess(len(skill_lines), 500)
        self.assertTrue((SKILL_ROOT / "references" / "ats-reality.md").exists())
        self.assertTrue((SKILL_ROOT / "references" / "interaction-flow.md").exists())
        self.assertTrue((SKILL_ROOT / "references" / "domains" / "catalog.json").exists())
        self.assertTrue((SKILL_ROOT / "schemas" / "career-record.schema.json").exists())
        self.assertTrue((SKILL_ROOT / "schemas" / "intake-session.schema.json").exists())


if __name__ == "__main__":
    unittest.main()
