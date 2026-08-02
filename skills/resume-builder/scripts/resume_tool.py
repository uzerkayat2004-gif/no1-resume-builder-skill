#!/usr/bin/env python3
"""Local validator, renderer, extractor, verifier, and interview scaffold.

This module deliberately performs no network calls. It treats the JSON career
record as the source of truth and blocks material claims that are unsupported
or marked unverified unless a record explicitly opts into an unverified draft.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import sys
import tempfile
import uuid
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from xml.sax.saxutils import escape


DATE_RE = re.compile(r"^(?:present|\d{4}(?:-\d{2}(?:-\d{2})?)?)?$", re.I)
PLACEHOLDER_RE = re.compile(
    r"\b(?:todo|tbd|xxx|placeholder)\b|\[(?:insert|add|confirm)[^\]]*\]",
    re.I,
)
NUMBER_RE = re.compile(
    r"(?<!\w)(?:[$€£₹]\s*)?\d+(?:[.,]\d+)?(?:\s?(?:%|[kKmMbB]))?"
)
WORD_RE = re.compile(r"[\w+#./-]+", re.UNICODE)
MATERIAL_STATUSES = {"verified", "user-confirmed"}
MAX_INPUT_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 50
MAX_EXTRACTED_CHARACTERS = 500_000
MAX_DOCX_ENTRIES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100
DRAFT_LABEL = "DRAFT — UNVERIFIED CLAIMS — NOT READY TO SUBMIT"


def load_record(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError("The career record must be a JSON object.")
    return data


def _all_strings(value: Any) -> Iterable[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, list):
        for item in value:
            yield from _all_strings(item)
    elif isinstance(value, dict):
        for item in value.values():
            yield from _all_strings(item)


def _claim_refs(record: dict[str, Any]) -> Iterable[tuple[str, list[str], str]]:
    basics = record.get("basics", {})
    identity = " | ".join(
        str(part)
        for part in (
            basics.get("name", ""),
            basics.get("label", ""),
            basics.get("location", ""),
        )
        if part
    )
    if identity:
        yield "basics.identity", basics.get("claimIds", []), identity
    if basics.get("summary"):
        yield (
            "basics.summary",
            basics.get("summaryClaimIds", []),
            basics["summary"],
        )

    for index, skill in enumerate(record.get("skills", [])):
        text = f"{skill.get('name', '')}: {', '.join(skill.get('keywords', []))}"
        yield f"skills[{index}]", skill.get("claimIds", []), text

    for section in ("work", "projects"):
        for index, item in enumerate(record.get(section, [])):
            item_text = " | ".join(
                part
                for part in (
                    item.get("position", ""),
                    item.get("name", ""),
                    item.get("location", ""),
                    _date_range(item),
                    item.get("description", ""),
                    item.get("summary", ""),
                )
                if part
            )
            if item.get("claimIds"):
                yield (
                    f"{section}[{index}]",
                    item.get("claimIds", []),
                    item_text,
                )
            for h_index, highlight in enumerate(item.get("highlights", [])):
                yield (
                    f"{section}[{index}].highlights[{h_index}]",
                    highlight.get("claimIds", []),
                    highlight.get("text", ""),
                )

    for index, item in enumerate(record.get("education", [])):
        text = " | ".join(
            part
            for part in (
                item.get("studyType", ""),
                item.get("area", ""),
                item.get("institution", ""),
                item.get("location", ""),
                _date_range(item),
                item.get("score", ""),
            )
            if part
        )
        yield f"education[{index}]", item.get("claimIds", []), text
        for h_index, highlight in enumerate(item.get("highlights", [])):
            yield (
                f"education[{index}].highlights[{h_index}]",
                highlight.get("claimIds", []),
                highlight.get("text", ""),
            )

    for section, label_parts in (
        ("certificates", ("name", "issuer", "date")),
        ("awards", ("title", "awarder", "summary")),
        ("languages", ("language", "fluency")),
    ):
        for index, item in enumerate(record.get(section, [])):
            text = " | ".join(str(item.get(key, "")) for key in label_parts if item.get(key))
            yield f"{section}[{index}]", item.get("claimIds", []), text

    for index, requirement in enumerate(record.get("target", {}).get("requirements", [])):
        if requirement.get("status") in {"supported", "transferable"}:
            yield (
                f"target.requirements[{index}]",
                requirement.get("claimIds", []),
                requirement.get("text", ""),
            )


def _numbers(text: str) -> set[str]:
    return {
        re.sub(r"[\s,]", "", match.group(0)).lower()
        for match in NUMBER_RE.finditer(text)
    }


def _claim_digest(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text).strip().casefold()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def approval_digests(record: dict[str, Any]) -> dict[str, str]:
    """Bind every rendered material claim to the candidate-approved wording."""
    return {
        location: _claim_digest(text)
        for location, _, text in _claim_refs(record)
        if not location.startswith("target.") and text.strip()
    }


def _structure_errors(record: Any) -> list[str]:
    errors: list[str] = []
    if not isinstance(record, dict):
        return ["The career record must be a JSON object."]

    def require_object(value: Any, location: str) -> dict[str, Any]:
        if not isinstance(value, dict):
            errors.append(f"{location} must be an object.")
            return {}
        return value

    def require_array(value: Any, location: str) -> list[Any]:
        if not isinstance(value, list):
            errors.append(f"{location} must be an array.")
            return []
        return value

    def require_string(value: Any, location: str) -> None:
        if value is not None and not isinstance(value, str):
            errors.append(f"{location} must be a string.")

    def require_string_array(value: Any, location: str) -> list[Any]:
        items = require_array(value, location)
        for index, item in enumerate(items):
            if not isinstance(item, str):
                errors.append(f"{location}[{index}] must be a string.")
        return items

    def check_fields(item: dict[str, Any], location: str, fields: Iterable[str]) -> None:
        for field in fields:
            if field in item:
                require_string(item[field], f"{location}.{field}")

    basics = require_object(record.get("basics"), "basics")
    check_fields(
        basics,
        "basics",
        ("name", "label", "email", "phone", "location", "url", "summary"),
    )
    for index, profile in enumerate(require_array(basics.get("profiles", []), "basics.profiles")):
        profile_item = require_object(profile, f"basics.profiles[{index}]")
        check_fields(profile_item, f"basics.profiles[{index}]", ("network", "url"))
    require_string_array(basics.get("claimIds", []), "basics.claimIds")
    require_string_array(basics.get("summaryClaimIds", []), "basics.summaryClaimIds")

    for section in (
        "skills",
        "work",
        "projects",
        "education",
        "certificates",
        "awards",
        "languages",
        "evidence",
    ):
        for index, raw_item in enumerate(require_array(record.get(section, []), section)):
            item = require_object(raw_item, f"{section}[{index}]")
            require_string_array(item.get("claimIds", []), f"{section}[{index}].claimIds")
            field_map = {
                "skills": ("name",),
                "work": ("name", "position", "location", "startDate", "endDate", "summary"),
                "projects": ("name", "description", "url", "startDate", "endDate"),
                "education": (
                    "institution",
                    "area",
                    "studyType",
                    "location",
                    "startDate",
                    "endDate",
                    "score",
                ),
                "certificates": ("name", "issuer", "date", "url"),
                "awards": ("title", "awarder", "date", "summary"),
                "languages": ("language", "fluency"),
                "evidence": ("id", "claim", "sourceType", "sourceStatus", "sourceNote"),
            }
            check_fields(item, f"{section}[{index}]", field_map[section])
            if "sensitive" in item and not isinstance(item["sensitive"], bool):
                errors.append(f"{section}[{index}].sensitive must be a boolean.")
            if section == "skills":
                require_string_array(item.get("keywords", []), f"skills[{index}].keywords")
            if section in {"work", "projects", "education"}:
                for h_index, raw_highlight in enumerate(
                    require_array(item.get("highlights", []), f"{section}[{index}].highlights")
                ):
                    highlight = require_object(
                        raw_highlight, f"{section}[{index}].highlights[{h_index}]"
                    )
                    check_fields(
                        highlight,
                        f"{section}[{index}].highlights[{h_index}]",
                        ("text",),
                    )
                    require_string_array(
                        highlight.get("claimIds", []),
                        f"{section}[{index}].highlights[{h_index}].claimIds",
                    )

    target = require_object(record.get("target", {}), "target")
    check_fields(
        target, "target", ("role", "level", "company", "jobDescription")
    )
    for index, raw_requirement in enumerate(
        require_array(target.get("requirements", []), "target.requirements")
    ):
        requirement = require_object(raw_requirement, f"target.requirements[{index}]")
        check_fields(
            requirement, f"target.requirements[{index}]", ("text", "status")
        )
        require_string_array(
            requirement.get("claimIds", []), f"target.requirements[{index}].claimIds"
        )
    meta = require_object(record.get("meta", {}), "meta")
    check_fields(meta, "meta", ("language", "region", "paperSize", "version", "lastModified"))
    if "allowUnverified" in meta and not isinstance(meta["allowUnverified"], bool):
        errors.append("meta.allowUnverified must be a boolean.")
    digests = meta.get("approvedClaimDigests", {})
    if not isinstance(digests, dict) or not all(
        isinstance(key, str) and isinstance(value, str)
        for key, value in getattr(digests, "items", lambda: [])()
    ):
        errors.append("meta.approvedClaimDigests must be an object of string digests.")
    return errors


def validate_record(
    record: dict[str, Any], *, require_approval: bool = True
) -> dict[str, Any]:
    errors = _structure_errors(record)
    warnings: list[str] = []
    if errors:
        return {
            "valid": False,
            "errors": errors,
            "warnings": warnings,
            "evidence": {"total": 0, "used": 0, "verified_or_confirmed": 0, "unverified": 0},
        }
    basics = record.get("basics")
    if not isinstance(basics, dict):
        errors.append("Missing basics object.")
        basics = {}
    if not str(basics.get("name", "")).strip():
        errors.append("basics.name is required.")

    contact_values = [
        basics.get("email"),
        basics.get("phone"),
        basics.get("url"),
        *(profile.get("url") for profile in basics.get("profiles", []) if isinstance(profile, dict)),
    ]
    if not any(str(value or "").strip() for value in contact_values):
        errors.append("At least one contact method or professional profile is required.")

    evidence_items = record.get("evidence")
    if not isinstance(evidence_items, list):
        errors.append("evidence must be an array.")
        evidence_items = []

    evidence: dict[str, dict[str, Any]] = {}
    for index, item in enumerate(evidence_items):
        if not isinstance(item, dict):
            errors.append(f"evidence[{index}] must be an object.")
            continue
        evidence_id = str(item.get("id", "")).strip()
        if not evidence_id:
            errors.append(f"evidence[{index}].id is required.")
            continue
        if evidence_id in evidence:
            errors.append(f"Duplicate evidence id: {evidence_id}.")
        evidence[evidence_id] = item
        if item.get("sourceStatus") not in {"verified", "user-confirmed", "unverified"}:
            errors.append(f"Evidence {evidence_id} has an invalid sourceStatus.")
        if not str(item.get("claim", "")).strip():
            errors.append(f"Evidence {evidence_id} has no claim text.")

    meta = record.get("meta", {})
    allow_unverified = bool(meta.get("allowUnverified", False))
    approved_digests = meta.get("approvedClaimDigests", {})
    used_evidence: set[str] = set()
    for location, claim_ids, text in _claim_refs(record):
        if not text.strip():
            warnings.append(f"{location} has no text.")
        if not claim_ids:
            errors.append(f"{location} has no evidence claimIds.")
            continue
        cited_numbers: set[str] = set()
        for claim_id in claim_ids:
            used_evidence.add(claim_id)
            item = evidence.get(claim_id)
            if item is None:
                errors.append(f"{location} references unknown evidence id {claim_id}.")
                continue
            status = item.get("sourceStatus")
            if status not in MATERIAL_STATUSES and not allow_unverified:
                errors.append(
                    f"{location} uses unverified evidence {claim_id}; confirm it or mark this as an unverified draft."
                )
            cited_numbers.update(_numbers(str(item.get("claim", ""))))
        unsupported_numbers = _numbers(text) - cited_numbers
        if unsupported_numbers:
            errors.append(
                f"{location} contains numeric detail not found in its cited evidence: "
                + ", ".join(sorted(unsupported_numbers))
            )
        if (
            require_approval
            and not allow_unverified
            and not location.startswith("target.")
            and approved_digests.get(location) != _claim_digest(text)
        ):
            errors.append(
                f"{location} changed or was not explicitly approved by the candidate; "
                "run the approval step after fact review."
            )

    if allow_unverified:
        warnings.append(
            "This is an unverified draft. Every output will be visibly marked and must not be submitted."
        )

    for section in ("work", "projects"):
        for index, item in enumerate(record.get(section, [])):
            if section == "work":
                for key in ("name", "position", "startDate"):
                    if not str(item.get(key, "")).strip():
                        errors.append(f"{section}[{index}].{key} is required.")
            if section == "projects" and not str(item.get("name", "")).strip():
                errors.append(f"projects[{index}].name is required.")
            for date_key in ("startDate", "endDate"):
                value = str(item.get(date_key, ""))
                if value and not DATE_RE.match(value):
                    errors.append(f"{section}[{index}].{date_key} has invalid date {value}.")

    for index, item in enumerate(record.get("education", [])):
        if not str(item.get("institution", "")).strip():
            errors.append(f"education[{index}].institution is required.")
        if not str(item.get("area", "")).strip():
            errors.append(f"education[{index}].area is required.")
        for date_key in ("startDate", "endDate"):
            value = str(item.get(date_key, ""))
            if value and not DATE_RE.match(value):
                errors.append(f"education[{index}].{date_key} has invalid date {value}.")

    placeholder_hits = sorted(
        {
            match.group(0)
            for text in _all_strings(record)
            for match in PLACEHOLDER_RE.finditer(text)
        }
    )
    if placeholder_hits:
        errors.append("Placeholder text remains: " + ", ".join(placeholder_hits))

    unused = sorted(set(evidence) - used_evidence)
    if unused:
        warnings.append(
            f"{len(unused)} evidence item(s) are unused in this variant: " + ", ".join(unused)
        )

    requirements = record.get("target", {}).get("requirements", [])
    gaps = [
        item.get("text", "")
        for item in requirements
        if item.get("status") in {"unsupported", "unclear"}
    ]
    if gaps:
        warnings.append(
            f"{len(gaps)} target requirement(s) are unsupported or unclear: "
            + " | ".join(gaps)
        )

    return {
        "valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "evidence": {
            "total": len(evidence),
            "used": len(used_evidence),
            "verified_or_confirmed": sum(
                1 for item in evidence.values() if item.get("sourceStatus") in MATERIAL_STATUSES
            ),
            "unverified": sum(
                1 for item in evidence.values() if item.get("sourceStatus") == "unverified"
            ),
        },
    }


def approve_record(record: dict[str, Any]) -> dict[str, Any]:
    """Create no-drift approval digests after the candidate confirms every fact."""
    candidate = copy.deepcopy(record)
    candidate.setdefault("meta", {})["allowUnverified"] = False
    report = validate_record(candidate, require_approval=False)
    if not report["valid"]:
        raise ValueError(
            "Career record failed pre-approval validation:\n- "
            + "\n- ".join(report["errors"])
        )
    candidate["meta"]["approvedClaimDigests"] = approval_digests(candidate)
    return candidate


def _month_label(value: str) -> str:
    if not value:
        return ""
    if value.lower() == "present":
        return "Present"
    try:
        if len(value) == 4:
            return value
        if len(value) == 7:
            return datetime.strptime(value, "%Y-%m").strftime("%b %Y")
        if len(value) == 10:
            return datetime.strptime(value, "%Y-%m-%d").strftime("%b %Y")
    except ValueError:
        return value
    return value


def _date_range(item: dict[str, Any]) -> str:
    start = _month_label(str(item.get("startDate", "")))
    end = _month_label(str(item.get("endDate", "")))
    if start and end:
        return f"{start} - {end}"
    return start or end


def _contact_line(basics: dict[str, Any]) -> str:
    parts = [
        basics.get("email"),
        basics.get("phone"),
        basics.get("location"),
        basics.get("url"),
    ]
    parts.extend(
        profile.get("url")
        for profile in basics.get("profiles", [])
        if isinstance(profile, dict)
    )
    return " | ".join(str(part).strip() for part in parts if str(part or "").strip())


def document_sections(record: dict[str, Any]) -> list[tuple[str, list[dict[str, Any]]]]:
    sections: list[tuple[str, list[dict[str, Any]]]] = []
    basics = record.get("basics", {})
    if basics.get("summary"):
        sections.append(("SUMMARY", [{"kind": "paragraph", "text": basics["summary"]}]))

    if record.get("skills"):
        items = []
        for skill in record["skills"]:
            prefix = f"{skill.get('name')}: " if skill.get("name") else ""
            items.append(
                {
                    "kind": "paragraph",
                    "text": prefix + ", ".join(skill.get("keywords", [])),
                }
            )
        sections.append(("SKILLS", items))

    if record.get("work"):
        items = []
        for role in record["work"]:
            items.append(
                {
                    "kind": "entry",
                    "title": " | ".join(
                        part for part in (role.get("position"), role.get("name")) if part
                    ),
                    "meta": " | ".join(
                        part
                        for part in (role.get("location"), _date_range(role))
                        if part
                    ),
                    "summary": role.get("summary", ""),
                    "bullets": [item["text"] for item in role.get("highlights", [])],
                }
            )
        sections.append(("EXPERIENCE", items))

    if record.get("projects"):
        items = []
        for project in record["projects"]:
            title_parts = [project.get("name")]
            if project.get("url"):
                title_parts.append(project.get("url"))
            items.append(
                {
                    "kind": "entry",
                    "title": " | ".join(part for part in title_parts if part),
                    "meta": _date_range(project),
                    "summary": project.get("description", ""),
                    "bullets": [item["text"] for item in project.get("highlights", [])],
                }
            )
        sections.append(("PROJECTS", items))

    if record.get("education"):
        items = []
        for education in record["education"]:
            degree = ", ".join(
                part
                for part in (education.get("studyType"), education.get("area"))
                if part
            )
            title = " | ".join(
                part for part in (degree, education.get("institution")) if part
            )
            meta = " | ".join(
                part
                for part in (
                    education.get("location"),
                    _date_range(education),
                    education.get("score"),
                )
                if part
            )
            items.append(
                {
                    "kind": "entry",
                    "title": title,
                    "meta": meta,
                    "bullets": [
                        highlight["text"]
                        for highlight in education.get("highlights", [])
                    ],
                }
            )
        sections.append(("EDUCATION", items))

    if record.get("certificates"):
        items = []
        for certificate in record["certificates"]:
            items.append(
                {
                    "kind": "entry",
                    "title": " | ".join(
                        part
                        for part in (
                            certificate.get("name"),
                            certificate.get("issuer"),
                        )
                        if part
                    ),
                    "meta": _month_label(str(certificate.get("date", ""))),
                    "summary": certificate.get("url", ""),
                    "bullets": [],
                }
            )
        sections.append(("CERTIFICATIONS", items))

    if record.get("awards"):
        items = []
        for award in record["awards"]:
            items.append(
                {
                    "kind": "entry",
                    "title": " | ".join(
                        part for part in (award.get("title"), award.get("awarder")) if part
                    ),
                    "meta": _month_label(str(award.get("date", ""))),
                    "summary": award.get("summary", ""),
                    "bullets": [],
                }
            )
        sections.append(("AWARDS", items))

    if record.get("languages"):
        text = ", ".join(
            " - ".join(
                part for part in (item.get("language"), item.get("fluency")) if part
            )
            for item in record["languages"]
        )
        sections.append(("LANGUAGES", [{"kind": "paragraph", "text": text}]))

    return sections


def build_plain_text(record: dict[str, Any]) -> str:
    basics = record.get("basics", {})
    lines = []
    if record.get("meta", {}).get("allowUnverified", False):
        lines.extend([DRAFT_LABEL, ""])
    lines.append(basics.get("name", ""))
    if basics.get("label"):
        lines.append(basics["label"])
    contact = _contact_line(basics)
    if contact:
        lines.append(contact)
    for title, items in document_sections(record):
        lines.extend(["", title])
        for item in items:
            if item["kind"] == "paragraph":
                lines.append(item["text"])
                continue
            lines.append(item.get("title", ""))
            if item.get("meta"):
                lines.append(item["meta"])
            if item.get("summary"):
                lines.append(item["summary"])
            lines.extend(f"• {bullet}" for bullet in item.get("bullets", []))
    return "\n".join(line for line in lines if line is not None).strip() + "\n"


def build_markdown(record: dict[str, Any]) -> str:
    basics = record.get("basics", {})
    lines = []
    if record.get("meta", {}).get("allowUnverified", False):
        lines.extend([f"> **{DRAFT_LABEL}**", ""])
    lines.append(f"# {basics.get('name', '')}")
    if basics.get("label"):
        lines.append(f"**{basics['label']}**")
    contact = _contact_line(basics)
    if contact:
        lines.append(contact)
    for title, items in document_sections(record):
        lines.extend(["", f"## {title.title()}"])
        for item in items:
            if item["kind"] == "paragraph":
                lines.append(item["text"])
                continue
            lines.append(f"### {item.get('title', '')}")
            if item.get("meta"):
                lines.append(f"*{item['meta']}*")
            if item.get("summary"):
                lines.append(item["summary"])
            lines.extend(f"- {bullet}" for bullet in item.get("bullets", []))
    return "\n".join(lines).strip() + "\n"


def _docx_imports() -> dict[str, Any]:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError(
            "DOCX rendering requires python-docx. Install requirements.txt."
        ) from exc
    return locals()


def _set_font(run: Any, name: str, size: float, bold: bool = False, italic: bool = False) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _add_style(document: Any, name: str, size: float, **kwargs: Any) -> Any:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    if name in document.styles:
        style = document.styles[name]
    else:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = kwargs.get("bold", False)
    style.font.italic = kwargs.get("italic", False)
    style.font.color.rgb = RGBColor.from_string(kwargs.get("color", "000000"))
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(kwargs.get("before", 0))
    paragraph.space_after = Pt(kwargs.get("after", 0))
    paragraph.line_spacing = kwargs.get("line_spacing", 1.0)
    paragraph.keep_with_next = kwargs.get("keep_with_next", False)
    if kwargs.get("center", False):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "left_indent" in kwargs:
        paragraph.left_indent = Inches(kwargs["left_indent"])
    if "first_line_indent" in kwargs:
        paragraph.first_line_indent = Inches(kwargs["first_line_indent"])
    return style


def _create_bullet_numbering(document: Any) -> int:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
        if item.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(item.get(qn("w:numId")))
        for item in numbering.findall(qn("w:num"))
        if item.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)

    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "360")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "360")
    ind.set(qn("w:hanging"), "180")
    ppr.append(ind)
    level.append(ppr)
    abstract.append(level)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)

    numbering.append(abstract)
    numbering.append(num)
    return num_id


def _apply_bullet_numbering(paragraph: Any, num_id: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def render_docx(record: dict[str, Any], output: Path) -> None:
    deps = _docx_imports()
    Document = deps["Document"]
    Cm = deps["Cm"]
    Inches = deps["Inches"]

    document = Document()
    section = document.sections[0]
    paper_size = record.get("meta", {}).get("paperSize", "letter").lower()
    if paper_size == "a4":
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, attr, Inches(0.70))
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = deps["Pt"](10.5)
    normal._element.get_or_add_rPr().rFonts.set(deps["qn"]("w:ascii"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(deps["qn"]("w:hAnsi"), "Arial")
    normal.paragraph_format.space_before = deps["Pt"](0)
    normal.paragraph_format.space_after = deps["Pt"](0)
    normal.paragraph_format.line_spacing = 1.0

    _add_style(document, "Resume Name", 20, bold=True, center=True, after=1)
    _add_style(document, "Resume Draft", 10, bold=True, center=True, after=5, color="B00020")
    _add_style(document, "Resume Label", 11, bold=True, center=True, after=1)
    _add_style(document, "Resume Contact", 9, center=True, after=7)
    _add_style(
        document,
        "Resume Section",
        10.5,
        bold=True,
        before=7,
        after=2,
        keep_with_next=True,
    )
    _add_style(
        document,
        "Resume Role",
        10.5,
        bold=True,
        before=3,
        after=0,
        keep_with_next=True,
    )
    _add_style(
        document,
        "Resume Meta",
        9.5,
        italic=True,
        after=1,
        keep_with_next=True,
        color="333333",
    )
    _add_style(document, "Resume Body", 10.5, after=2, line_spacing=1.0)
    _add_style(
        document,
        "Resume Bullet",
        10.2,
        after=1.5,
        line_spacing=1.0,
        left_indent=0.25,
        first_line_indent=-0.125,
    )
    bullet_num_id = _create_bullet_numbering(document)

    basics = record.get("basics", {})
    if record.get("meta", {}).get("allowUnverified", False):
        document.add_paragraph(DRAFT_LABEL, style="Resume Draft")
    document.add_paragraph(basics.get("name", ""), style="Resume Name")
    if basics.get("label"):
        document.add_paragraph(basics["label"], style="Resume Label")
    contact = _contact_line(basics)
    if contact:
        document.add_paragraph(contact, style="Resume Contact")

    for title, items in document_sections(record):
        document.add_paragraph(title, style="Resume Section")
        for item in items:
            if item["kind"] == "paragraph":
                document.add_paragraph(item["text"], style="Resume Body")
                continue
            document.add_paragraph(item.get("title", ""), style="Resume Role")
            if item.get("meta"):
                document.add_paragraph(item["meta"], style="Resume Meta")
            if item.get("summary"):
                document.add_paragraph(item["summary"], style="Resume Body")
            for bullet in item.get("bullets", []):
                paragraph = document.add_paragraph(style="Resume Bullet")
                _apply_bullet_numbering(paragraph, bullet_num_id)
                paragraph.add_run(bullet)

    properties = document.core_properties
    properties.title = f"{basics.get('name', 'Candidate')} - Resume"
    properties.subject = ""
    properties.author = ""
    properties.keywords = ""
    properties.comments = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    scrub_docx_metadata(output)


def scrub_docx_metadata(output: Path) -> None:
    """Remove author/time/revision metadata without changing visible content."""
    temporary = output.with_suffix(output.suffix + ".tmp")
    try:
        with zipfile.ZipFile(output, "r") as source, zipfile.ZipFile(
            temporary, "w", compression=zipfile.ZIP_DEFLATED
        ) as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    for tag in (
                        b"dc:creator",
                        b"cp:lastModifiedBy",
                        b"cp:revision",
                        b"dcterms:created",
                        b"dcterms:modified",
                    ):
                        pattern = (
                            rb"<" + tag + rb"\b[^>]*(?:/>|>.*?</" + tag + rb"\s*>)"
                        )
                        data = re.sub(pattern, b"", data, flags=re.DOTALL)
                if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                    data = re.sub(
                        rb"\s+w:rsid[A-Za-z0-9]*=\"[^\"]*\"", b"", data
                    )
                    data = re.sub(
                        rb"<w:rsids\b[^>]*>.*?</w:rsids\s*>",
                        b"",
                        data,
                        flags=re.DOTALL,
                    )
                    data = re.sub(rb"<w:rsid[A-Za-z0-9]*\b[^>]*/>", b"", data)
                destination.writestr(info, data)
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()


def _pdf_fonts() -> dict[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = {
        "regular": [
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        ],
        "bold": [
            Path("C:/Windows/Fonts/arialbd.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
        ],
        "italic": [
            Path("C:/Windows/Fonts/ariali.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Italic.ttf"),
        ],
    }
    selected = {
        key: next((path for path in paths if path.exists()), None)
        for key, paths in candidates.items()
    }
    if all(selected.values()):
        pdfmetrics.registerFont(TTFont("ResumeSans", str(selected["regular"])))
        pdfmetrics.registerFont(TTFont("ResumeSans-Bold", str(selected["bold"])))
        pdfmetrics.registerFont(TTFont("ResumeSans-Italic", str(selected["italic"])))
        return {
            "regular": "ResumeSans",
            "bold": "ResumeSans-Bold",
            "italic": "ResumeSans-Italic",
        }
    return {
        "regular": "Helvetica",
        "bold": "Helvetica-Bold",
        "italic": "Helvetica-Oblique",
    }


def render_pdf(record: dict[str, Any], output: Path) -> None:
    try:
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, LETTER
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError(
            "PDF rendering requires reportlab. Install requirements.txt."
        ) from exc

    fonts = _pdf_fonts()
    page_size = (
        A4
        if record.get("meta", {}).get("paperSize", "letter").lower() == "a4"
        else LETTER
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    basics = record.get("basics", {})
    doc = SimpleDocTemplate(
        str(output),
        pagesize=page_size,
        rightMargin=0.70 * inch,
        leftMargin=0.70 * inch,
        topMargin=0.66 * inch,
        bottomMargin=0.66 * inch,
        title=f"{basics.get('name', 'Candidate')} - Resume",
        author="",
        subject="",
    )
    styles = {
        "name": ParagraphStyle(
            "ResumeName",
            fontName=fonts["bold"],
            fontSize=20,
            leading=22,
            alignment=TA_CENTER,
            spaceAfter=1,
        ),
        "draft": ParagraphStyle(
            "ResumeDraft",
            fontName=fonts["bold"],
            fontSize=9.5,
            leading=11,
            textColor="#B00020",
            alignment=TA_CENTER,
            spaceAfter=5,
        ),
        "label": ParagraphStyle(
            "ResumeLabel",
            fontName=fonts["bold"],
            fontSize=11,
            leading=13,
            alignment=TA_CENTER,
            spaceAfter=1,
        ),
        "contact": ParagraphStyle(
            "ResumeContact",
            fontName=fonts["regular"],
            fontSize=8.8,
            leading=10.5,
            alignment=TA_CENTER,
            spaceAfter=7,
        ),
        "section": ParagraphStyle(
            "ResumeSection",
            fontName=fonts["bold"],
            fontSize=10.5,
            leading=12,
            spaceBefore=7,
            spaceAfter=2,
            keepWithNext=True,
        ),
        "role": ParagraphStyle(
            "ResumeRole",
            fontName=fonts["bold"],
            fontSize=10.5,
            leading=12,
            spaceBefore=3,
            spaceAfter=0,
            keepWithNext=True,
        ),
        "meta": ParagraphStyle(
            "ResumeMeta",
            fontName=fonts["italic"],
            fontSize=9.3,
            leading=11,
            spaceAfter=1,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "ResumeBody",
            fontName=fonts["regular"],
            fontSize=10.3,
            leading=12,
            spaceAfter=2,
        ),
        "bullet": ParagraphStyle(
            "ResumeBullet",
            fontName=fonts["regular"],
            fontSize=10.1,
            leading=11.8,
            leftIndent=14,
            firstLineIndent=-8,
            spaceAfter=1.5,
        ),
    }
    story: list[Any] = []
    if record.get("meta", {}).get("allowUnverified", False):
        story.append(Paragraph(escape(DRAFT_LABEL), styles["draft"]))
    story.append(Paragraph(escape(str(basics.get("name", ""))), styles["name"]))
    if basics.get("label"):
        story.append(Paragraph(escape(str(basics["label"])), styles["label"]))
    contact = _contact_line(basics)
    if contact:
        story.append(Paragraph(escape(contact), styles["contact"]))

    for title, items in document_sections(record):
        story.append(Paragraph(escape(title), styles["section"]))
        for item in items:
            if item["kind"] == "paragraph":
                story.append(Paragraph(escape(item["text"]), styles["body"]))
                continue
            heading_parts = [Paragraph(escape(item.get("title", "")), styles["role"])]
            if item.get("meta"):
                heading_parts.append(Paragraph(escape(item["meta"]), styles["meta"]))
            story.append(KeepTogether(heading_parts))
            if item.get("summary"):
                story.append(Paragraph(escape(item["summary"]), styles["body"]))
            for bullet in item.get("bullets", []):
                story.append(
                    Paragraph(f"• {escape(str(bullet))}", styles["bullet"])
                )
    doc.build(story)
    scrub_pdf_metadata(output)


def scrub_pdf_metadata(output: Path) -> None:
    """Rewrite a generated PDF without timestamps, actions, or attachments."""
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise RuntimeError("PDF metadata scrubbing requires pypdf.") from exc
    reader = PdfReader(str(output))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    title = str((reader.metadata or {}).get("/Title", "Resume"))
    writer.add_metadata({"/Title": title, "/Author": "", "/Subject": ""})
    temporary = output.with_suffix(output.suffix + ".tmp")
    try:
        with temporary.open("wb") as handle:
            writer.write(handle)
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()


def _sanitize_basename(raw: str) -> str:
    value = re.sub(r"[^\w.-]+", "_", str(raw), flags=re.UNICODE)
    value = re.sub(r"_+", "_", value).strip("_.")[:120]
    reserved = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
    if value.upper() in reserved:
        value = f"Resume_{value}"
    return value or "Resume"


def safe_basename(record: dict[str, Any]) -> str:
    basics = record.get("basics", {})
    target = record.get("target", {})
    raw = "_".join(
        part
        for part in (
            basics.get("name", "Candidate"),
            target.get("role") or basics.get("label") or "Resume",
            "Resume",
        )
        if part
    )
    return _sanitize_basename(raw)


def render_all(
    record: dict[str, Any],
    outdir: Path,
    formats: list[str],
    basename: str | None = None,
) -> dict[str, Any]:
    validation = validate_record(record)
    if not validation["valid"]:
        raise ValueError(
            "Career record failed validation:\n- " + "\n- ".join(validation["errors"])
        )
    outdir.mkdir(parents=True, exist_ok=True)
    stem = _sanitize_basename(basename) if basename else safe_basename(record)
    is_draft = bool(record.get("meta", {}).get("allowUnverified", False))
    if is_draft and not stem.upper().endswith("_DRAFT"):
        stem += "_DRAFT"
    files: dict[str, str] = {}
    for format_name in formats:
        format_name = format_name.lower().strip()
        output = outdir / f"{stem}.{format_name}"
        if format_name == "docx":
            render_docx(record, output)
        elif format_name == "pdf":
            render_pdf(record, output)
        elif format_name == "txt":
            output.write_text(build_plain_text(record), encoding="utf-8")
        elif format_name == "md":
            output.write_text(build_markdown(record), encoding="utf-8")
        else:
            raise ValueError(f"Unsupported output format: {format_name}")
        files[format_name] = output.name

    manifest = {
        "recordValid": True,
        "draft": is_draft,
        "basename": stem,
        "formats": files,
        "warnings": validation["warnings"],
        "note": "These files were rendered from the evidence-backed career record. Employer application instructions override generic format guidance.",
    }
    manifest_path = outdir / f"{stem}.manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    manifest["manifest"] = manifest_path.name
    return manifest


def _check_input_file(document_path: Path) -> None:
    if not document_path.exists() or not document_path.is_file():
        raise ValueError(f"Document not found: {document_path}")
    size = document_path.stat().st_size
    if size > MAX_INPUT_BYTES:
        raise ValueError(
            f"Document is too large ({size} bytes); limit is {MAX_INPUT_BYTES} bytes."
        )


def _check_docx_archive(document_path: Path) -> None:
    try:
        with zipfile.ZipFile(document_path, "r") as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise ValueError("DOCX contains too many archive entries.")
            total_compressed = 0
            total_uncompressed = 0
            for entry in entries:
                normalized = entry.filename.replace("\\", "/")
                parts = Path(normalized).parts
                if normalized.startswith("/") or ".." in parts:
                    raise ValueError("DOCX contains an unsafe archive path.")
                if entry.flag_bits & 0x1:
                    raise ValueError("Encrypted DOCX archive entries are not supported.")
                lower = normalized.casefold()
                if lower.endswith("vbaproject.bin") or "/embeddings/" in f"/{lower}":
                    raise ValueError("Macro-enabled or embedded-object DOCX files are not accepted.")
                if lower.endswith((".xml", ".rels")):
                    prefix = archive.read(entry)[:8192].upper()
                    if b"<!DOCTYPE" in prefix or b"<!ENTITY" in prefix:
                        raise ValueError("DOCX XML declarations with entities are not accepted.")
                total_compressed += max(entry.compress_size, 1)
                total_uncompressed += entry.file_size
            if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX expands beyond the safe uncompressed-size limit.")
            if (
                total_uncompressed > 1_000_000
                and total_uncompressed / max(total_compressed, 1)
                > MAX_DOCX_COMPRESSION_RATIO
            ):
                raise ValueError("DOCX has a suspicious archive compression ratio.")
    except zipfile.BadZipFile as exc:
        raise ValueError("DOCX is not a valid Office archive.") from exc


def _bounded_text(text: str) -> str:
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError(
            "Extracted text exceeds the safe character limit "
            f"({MAX_EXTRACTED_CHARACTERS})."
        )
    return text.strip() + "\n" if text.strip() else ""


def extract_text(document_path: Path) -> str:
    _check_input_file(document_path)
    suffix = document_path.suffix.lower()
    if suffix == ".docx":
        _check_docx_archive(document_path)
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX extraction requires python-docx.") from exc
        document = Document(document_path)
        chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        return _bounded_text("\n".join(chunks))
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF extraction requires pypdf.") from exc
        reader = PdfReader(str(document_path))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"PDF has more than {MAX_PDF_PAGES} pages.")
        chunks: list[str] = []
        total = 0
        for page in reader.pages:
            chunk = page.extract_text() or ""
            total += len(chunk)
            if total > MAX_EXTRACTED_CHARACTERS:
                raise ValueError("PDF extracted text exceeds the safe character limit.")
            chunks.append(chunk)
        return _bounded_text("\n".join(chunks))
    if suffix in {".txt", ".md"}:
        return _bounded_text(document_path.read_text(encoding="utf-8"))
    raise ValueError(f"Unsupported document type: {suffix}")


def _domain_catalog_path() -> Path:
    return Path(__file__).resolve().parents[1] / "references" / "domains" / "catalog.json"


def load_domain_catalog(path: Path | None = None) -> dict[str, Any]:
    catalog_path = path or _domain_catalog_path()
    with catalog_path.open("r", encoding="utf-8") as handle:
        catalog = json.load(handle)
    if not isinstance(catalog.get("domains"), list) or not catalog.get("fallback"):
        raise ValueError(f"Invalid domain catalog: {catalog_path}")
    return catalog


def _keyword_match(text: str, keyword: str) -> bool:
    normalized_keyword = re.sub(r"\s+", " ", keyword.casefold().strip())
    if not normalized_keyword:
        return False
    pattern = r"(?<!\w)" + re.escape(normalized_keyword) + r"(?!\w)"
    return re.search(pattern, text) is not None


def route_domain(query: str, catalog_path: Path | None = None) -> dict[str, Any]:
    catalog = load_domain_catalog(catalog_path)
    normalized = re.sub(r"\s+", " ", query.casefold())
    matches: list[dict[str, Any]] = []
    for domain in catalog["domains"]:
        matched_keywords = [
            keyword
            for keyword in domain.get("keywords", [])
            if _keyword_match(normalized, keyword)
        ]
        score = sum(max(2, len(keyword.split()) * 2) for keyword in matched_keywords)
        if score:
            matches.append(
                {
                    "id": domain["id"],
                    "label": domain["label"],
                    "profile": domain["profile"],
                    "score": score,
                    "matchedKeywords": matched_keywords,
                    "evidencePrompts": domain.get("evidencePrompts", []),
                }
            )
    matches.sort(key=lambda item: (-item["score"], item["id"]))
    if not matches:
        fallback = catalog["fallback"]
        return {
            "selected": {
                "id": fallback["id"],
                "label": fallback["label"],
                "profile": fallback["profile"],
                "score": 0,
                "matchedKeywords": [],
                "evidencePrompts": fallback.get("evidencePrompts", []),
            },
            "confidence": "low",
            "needsConfirmation": True,
            "alternatives": [],
            "interpretation": "The domain router selects question guidance; it does not infer candidate skills.",
        }

    selected = matches[0]
    second_score = matches[1]["score"] if len(matches) > 1 else 0
    if selected["score"] >= 6 and (
        second_score == 0 or selected["score"] >= second_score * 1.5
    ):
        confidence = "high"
    elif selected["score"] >= 3:
        confidence = "medium"
    else:
        confidence = "low"
    return {
        "selected": selected,
        "confidence": confidence,
        "needsConfirmation": confidence != "high" or (
            len(matches) > 1 and selected["score"] == second_score
        ),
        "alternatives": matches[1:4],
        "interpretation": "The domain router selects question guidance; it does not infer candidate skills.",
    }


def _utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def summarize_imported_text(text: str) -> dict[str, Any]:
    """Return a privacy-preserving understanding report without copying source text."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    heading_names = {
        "summary": ("summary", "profile", "objective"),
        "experience": ("experience", "employment", "work history"),
        "education": ("education", "academic background"),
        "skills": ("skills", "technical skills", "competencies"),
        "projects": ("projects", "selected projects"),
        "certifications": ("certifications", "licenses"),
    }
    normalized_lines = {re.sub(r"[^a-z ]", "", line.casefold()).strip() for line in lines}
    detected_sections = [
        label
        for label, aliases in heading_names.items()
        if any(alias in normalized_lines for alias in aliases)
    ]
    sensitive_patterns = {
        "date of birth": r"\b(?:date of birth|dob)\b",
        "marital status": r"\bmarital status\b",
        "religion": r"\breligion\b",
        "government identifier": r"\b(?:passport|aadhaar|social security|ssn|national id)\b",
        "full street address": r"\b(?:street|road|avenue|lane|apartment|flat)\b",
    }
    sensitive_categories = [
        label for label, pattern in sensitive_patterns.items() if re.search(pattern, text, re.I)
    ]
    instruction_signals = len(
        re.findall(
            r"\b(?:ignore (?:all |the )?(?:previous|prior) instructions|system prompt|"
            r"run (?:this )?command|upload (?:the )?file|send (?:this|the data)|"
            r"click (?:this|the )?link)\b",
            text,
            re.I,
        )
    )
    date_signals = len(
        re.findall(
            r"\b(?:19|20)\d{2}(?:[-/]\d{1,2})?(?:\s*(?:-|–|to)\s*(?:present|(?:19|20)\d{2}(?:[-/]\d{1,2})?))?\b",
            text,
            re.I,
        )
    )
    warnings: list[str] = []
    if not text.strip():
        warnings.append("No selectable text was extracted; the file may be scanned.")
    if not detected_sections and text.strip():
        warnings.append("No standard resume section headings were confidently detected.")
    if sensitive_categories:
        warnings.append("Potentially sensitive categories were detected; review before reuse.")
    if instruction_signals:
        warnings.append(
            "Instruction-like text was detected and will be treated only as untrusted document data."
        )
    return {
        "nonEmptyLines": len(lines),
        "detectedSections": detected_sections,
        "dateSignals": date_signals,
        "potentialSensitiveCategories": sensitive_categories,
        "instructionLikeSignals": instruction_signals,
        "warnings": warnings,
    }


def new_intake_session(
    journey: str,
    source: Path | None = None,
    role: str = "",
    region: str = "",
    domain: str = "",
    job_description: str = "",
    include_extracted_text: bool = False,
) -> dict[str, Any]:
    if journey not in {"existing-resume", "scratch"}:
        raise ValueError("journey must be existing-resume or scratch.")
    if source and not source.exists():
        raise ValueError(f"Source file not found: {source}")

    sources: list[dict[str, Any]] = []
    extracted = ""
    if source:
        extracted = extract_text(source)
        sources.append(
            {
                "path": source.name,
                "fileName": source.name,
                "fileType": source.suffix.lower().lstrip("."),
                "sha256": _file_digest(source),
                "extractedCharacters": len(extracted.strip()),
                "textSelectable": bool(extracted.strip()),
                "extractedText": extracted if include_extracted_text else "",
                "extractedTextStored": include_extracted_text,
                "importSummary": summarize_imported_text(extracted),
                "reviewedByUser": False,
            }
        )

    route_input = " ".join(
        value
        for value in (
            role,
            job_description,
            extracted[:4000] if not role and not job_description else "",
        )
        if value
    )
    if domain:
        catalog = load_domain_catalog()
        selected = next(
            (item for item in catalog["domains"] if item["id"] == domain),
            None,
        )
        if selected is None and domain != catalog["fallback"]["id"]:
            raise ValueError(f"Unknown domain id: {domain}")
        selected = selected or catalog["fallback"]
        domain_result = {
            "selected": {
                "id": selected["id"],
                "label": selected["label"],
                "profile": selected["profile"],
                "score": None,
                "matchedKeywords": [],
                "evidencePrompts": selected.get("evidencePrompts", []),
            },
            "confidence": "user-selected",
            "needsConfirmation": False,
            "alternatives": [],
            "interpretation": "The domain was explicitly selected.",
        }
    else:
        domain_result = route_domain(route_input)

    now = _utc_now()
    session: dict[str, Any] = {
        "version": "1.2.0",
        "sessionId": str(uuid.uuid4()),
        "journey": journey,
        "status": "intake",
        "createdAt": now,
        "updatedAt": now,
        "privacyNotice": (
            "This session contains personal career data. Full extracted text is omitted by default. "
            "Keep the file private, do not commit it, and check your AI platform's data controls."
        ),
        "target": {
            "role": role,
            "region": region,
            "jobDescription": job_description,
            "domain": domain_result["selected"],
            "domainConfidence": domain_result["confidence"],
            "domainNeedsConfirmation": domain_result["needsConfirmation"],
        },
        "candidate": {
            "name": "",
            "email": "",
            "phone": "",
            "location": "",
            "profiles": [],
        },
        "sources": sources,
        "answers": {
            "sourceReviewed": False,
            "timeline": [],
            "timelineConfirmed": False,
            "evidenceStories": [],
            "education": [],
            "certifications": [],
            "skills": [],
            "projects": [],
            "languages": [],
            "finalReviewConfirmed": False,
        },
        "notes": [],
        "askedQuestionIds": [],
        "nextQuestions": [],
    }
    session["nextQuestions"] = next_intake_questions(session)
    return session


def intake_progress(session: dict[str, Any]) -> dict[str, str]:
    target = session.get("target", {})
    candidate = session.get("candidate", {})
    answers = session.get("answers", {})
    sources = session.get("sources", [])
    journey = session.get("journey")

    target_values = [target.get("role"), target.get("region")]
    target_status = (
        "complete"
        if all(target_values)
        else "partial"
        if any(target_values)
        else "pending"
    )
    contact_values = [
        candidate.get("email"),
        candidate.get("phone"),
        *(profile.get("url") for profile in candidate.get("profiles", []) if isinstance(profile, dict)),
    ]
    contact_status = (
        "complete"
        if candidate.get("name") and any(contact_values)
        else "partial"
        if candidate.get("name") or any(contact_values)
        else "pending"
    )
    if journey == "scratch":
        source_status = "not-applicable"
    elif not sources:
        source_status = "pending"
    elif answers.get("sourceReviewed"):
        source_status = "complete"
    else:
        source_status = "needs-confirmation"

    timeline = answers.get("timeline", [])
    timeline_status = (
        "complete"
        if timeline and answers.get("timelineConfirmed")
        else "partial"
        if timeline or sources
        else "pending"
    )
    stories = answers.get("evidenceStories", [])
    evidence_status = "complete" if len(stories) >= 2 else "partial" if stories else "pending"
    qualifications = (
        answers.get("education", [])
        + answers.get("certifications", [])
        + answers.get("skills", [])
        + answers.get("projects", [])
    )
    qualifications_status = "complete" if qualifications else "pending"
    review_status = "complete" if answers.get("finalReviewConfirmed") else "pending"
    return {
        "source": source_status,
        "target": target_status,
        "contact": contact_status,
        "timeline": timeline_status,
        "evidence": evidence_status,
        "qualifications": qualifications_status,
        "review": review_status,
    }


def _domain_prompts(session: dict[str, Any]) -> list[str]:
    selected = session.get("target", {}).get("domain", {})
    prompts = selected.get("evidencePrompts", [])
    if prompts:
        return prompts
    catalog = load_domain_catalog()
    return catalog["fallback"].get("evidencePrompts", [])


def next_intake_questions(session: dict[str, Any], limit: int = 3) -> list[dict[str, str]]:
    questions: list[dict[str, str]] = []
    journey = session.get("journey")
    target = session.get("target", {})
    candidate = session.get("candidate", {})
    answers = session.get("answers", {})
    sources = session.get("sources", [])
    asked = set(session.get("askedQuestionIds", []))

    def add(question_id: str, question: str, why: str = "") -> None:
        if len(questions) < limit and question_id not in asked:
            payload = {"id": question_id, "question": question}
            if why:
                payload["why"] = why
            questions.append(payload)

    if journey == "existing-resume" and not sources:
        add(
            "source.upload",
            "Please upload your current or old resume as PDF, DOCX, TXT, or Markdown.",
            "I will preserve it and first show what extracted before rewriting.",
        )
    if journey == "existing-resume" and sources and not answers.get("sourceReviewed"):
        extracted_count = sum(item.get("extractedCharacters", 0) for item in sources)
        add(
            "source.confirm",
            f"I extracted {extracted_count} characters. Is this your latest reliable career history, and are any sections outdated?",
        )
        add(
            "source.exclude",
            "Is there anything in the old resume you do not want reused, such as an address, photo, objective, or outdated skill?",
        )
    if not target.get("role"):
        add(
            "target.role",
            "What role or type of work are you targeting?",
        )
    if not target.get("region"):
        add(
            "target.region",
            "Which country or job market are you applying in?",
            "Formatting, privacy conventions, and document expectations vary by market.",
        )
    if target.get("domainNeedsConfirmation"):
        selected_label = target.get("domain", {}).get("label", "general")
        add(
            "target.domain",
            f"I routed this to “{selected_label}.” Does that fit your work, or should I use another domain?",
        )
    if not candidate.get("name"):
        add("contact.name", "What name should appear at the top of the resume?")
    contact_exists = candidate.get("email") or candidate.get("phone") or candidate.get("profiles")
    if not contact_exists:
        add(
            "contact.method",
            "Which safe contact detail should employers use: email, phone, or a professional profile?",
        )
    if not answers.get("timeline"):
        if journey == "scratch":
            add(
                "timeline.recent",
                "Tell me about your current or most recent work, project, training, volunteering, or responsibility: organization, role, and dates.",
            )
        elif answers.get("sourceReviewed"):
            add(
                "timeline.confirm",
                "Which imported role is most important for your target, and are its title and dates correct?",
            )
    if len(answers.get("evidenceStories", [])) < 2:
        for index, prompt in enumerate(_domain_prompts(session)):
            add(f"evidence.domain.{index}", prompt)
    qualifications = (
        answers.get("education", [])
        + answers.get("certifications", [])
        + answers.get("skills", [])
        + answers.get("projects", [])
    )
    if not qualifications:
        add(
            "qualifications.core",
            "Which education, licenses, certifications, tools, languages, or projects are most relevant to this target?",
        )
    if (
        intake_progress(session)["evidence"] == "complete"
        and not answers.get("finalReviewConfirmed")
    ):
        add(
            "review.confirm",
            "Before I render the final files, would you like to review the fact-check list and unsupported job requirements?",
        )
    return questions[:limit]


def intake_status_report(session: dict[str, Any]) -> dict[str, Any]:
    progress = intake_progress(session)
    complete = sum(
        1 for status in progress.values() if status in {"complete", "not-applicable"}
    )
    percent = round(complete / len(progress) * 100)
    return {
        "sessionId": session.get("sessionId"),
        "journey": session.get("journey"),
        "status": session.get("status"),
        "progress": progress,
        "completionPercent": percent,
        "selectedDomain": session.get("target", {}).get("domain", {}),
        "importSummaries": [
            item.get("importSummary", {}) for item in session.get("sources", [])
        ],
        "nextQuestions": next_intake_questions(session),
        "privacyNotice": session.get("privacyNotice"),
    }


def _tokens(text: str) -> list[str]:
    return [
        token.casefold()
        for token in WORD_RE.findall(text)
        if any(character.isalnum() for character in token)
    ]


def _normal_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def expected_anchors(record: dict[str, Any]) -> list[str]:
    anchors = [record.get("basics", {}).get("name", "")]
    anchors.extend(title for title, _ in document_sections(record))
    anchors.extend(item.get("name", "") for item in record.get("work", []))
    anchors.extend(item.get("name", "") for item in record.get("projects", []))
    anchors.extend(
        item.get("institution", "") for item in record.get("education", [])
    )
    return [anchor for anchor in anchors if anchor]


def verify_document(record: dict[str, Any], document_path: Path) -> dict[str, Any]:
    validation = validate_record(record)
    if not validation["valid"]:
        raise ValueError(
            "Career record failed validation:\n- " + "\n- ".join(validation["errors"])
        )
    expected = build_plain_text(record)
    actual = extract_text(document_path)
    expected_counter = Counter(_tokens(expected))
    actual_counter = Counter(_tokens(actual))
    matched = sum((expected_counter & actual_counter).values())
    expected_total = sum(expected_counter.values())
    recall = matched / expected_total if expected_total else 1.0
    actual_normal = _normal_text(actual)
    anchors = expected_anchors(record)
    missing_anchors = [
        anchor for anchor in anchors if _normal_text(anchor) not in actual_normal
    ]
    extracted_pages = None
    if document_path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        extracted_pages = len(PdfReader(str(document_path)).pages)
    report = {
        "document": document_path.name,
        "documentType": document_path.suffix.lower().lstrip("."),
        "extractableCharacters": len(actual.strip()),
        "expectedTokenCount": expected_total,
        "matchedTokenCount": matched,
        "tokenRecall": round(recall, 4),
        "anchorsChecked": anchors,
        "missingAnchors": missing_anchors,
        "pages": extracted_pages,
        "pass": recall >= 0.97 and not missing_anchors and bool(actual.strip()),
        "interpretation": (
            "This is a deterministic extraction check, not an ATS score or a prediction of employer behavior."
        ),
    }
    return report


def interview_markdown(record: dict[str, Any]) -> str:
    validation = validate_record(record)
    if not validation["valid"]:
        raise ValueError(
            "Career record failed validation:\n- " + "\n- ".join(validation["errors"])
        )
    role = (
        record.get("target", {}).get("role")
        or record.get("basics", {}).get("label")
        or "the target role"
    )
    lines = [
        f"# Interview practice - {role}",
        "",
        "Use these prompts to practice truthful, conversational answers. The claim IDs are evidence anchors, not text to recite.",
        "",
        "## Opening",
        "",
        f"1. Walk me through your background and why it is relevant to {role}.",
        "   - Build a 60-90 second answer from the summary and the two most relevant recent roles.",
        "",
        "## Screening questions",
        "",
        f"1. Why are you interested in {role}, and what evidence makes you a credible match?",
        "2. What are your availability, location, and work-authorization constraints? Answer only what is lawful and relevant in your market.",
        "",
        "## Requirement and technical/domain questions",
        "",
    ]
    requirements = record.get("target", {}).get("requirements", [])
    if requirements:
        for index, requirement in enumerate(requirements, start=1):
            status = requirement.get("status", "unclear")
            claims = ", ".join(requirement.get("claimIds", [])) or "none"
            lines.extend(
                [
                    f"{index}. The role requires: {requirement.get('text', '')}",
                    f"   - Evidence status: {status}",
                    f"   - Evidence anchors: {claims}",
                ]
            )
            if status in {"unsupported", "unclear"}:
                lines.append(
                    "   - Practice an honest gap response: acknowledge the gap, name adjacent evidence, and explain how you would close it."
                )
    else:
        lines.append(
            "1. Ask for the job description, then map each major requirement to verified evidence."
        )

    lines.extend(
        [
            "",
            "## Behavioral questions",
            "",
            "1. Tell me about a difficult problem you solved with incomplete information.",
            "2. Describe a disagreement with a stakeholder and how you handled it.",
            "3. Tell me about a mistake or failed approach and what changed afterward.",
            "   - Use STAR-R: Situation, Task, Action, Result, Reflection.",
            "",
            "## Resume deep dives",
            "",
        ]
    )
    question_number = 1
    for work_index, role_item in enumerate(record.get("work", [])[:3]):
        for highlight in role_item.get("highlights", [])[:2]:
            claims = ", ".join(highlight.get("claimIds", []))
            lines.extend(
                [
                    f"{question_number}. Walk me through this result: “{highlight.get('text', '')}”",
                    f"   - Evidence anchors: {claims}",
                    "   - Cover Situation, Task, Action, Result, and Reflection. Be ready to explain your exact contribution, trade-offs, and what you learned.",
                ]
            )
            question_number += 1

    lines.extend(
        [
            "",
            "## Honest gap responses",
            "",
            "- State the gap directly; do not imply experience you do not have.",
            "- Connect the closest verified, transferable evidence.",
            "- Explain a concrete learning or first-90-days plan without promising certainty.",
            "",
            "## Questions for the interviewer",
            "",
            "- What outcomes would define success in the first 90 days?",
            "- Which problems are most urgent for this role?",
            "- How are decisions made and work reviewed on this team?",
            "- What distinguishes strong performers here after one year?",
            "",
            "## Reusable follow-ups",
            "",
            "- What part did you personally own?",
            "- What alternatives did you consider?",
            "- How did you measure the result?",
            "- What would you do differently?",
            "- Which stakeholder disagreed, and how did you handle it?",
            "- What evidence can you show or discuss without exposing confidential information?",
            "",
            "## Practice rule",
            "",
            "Do not memorize a fictional script. Practice the structure, facts, and evidence until you can explain them naturally.",
        ]
    )
    return "\n".join(lines).strip() + "\n"


def _write_or_print(payload: Any, output: Path | None = None) -> None:
    text = (
        payload
        if isinstance(payload, str)
        else json.dumps(payload, indent=2, ensure_ascii=False) + "\n"
    )
    if output:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(text, encoding="utf-8")
    else:
        print(text, end="")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Evidence-backed resume validator and renderer."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    start = subparsers.add_parser(
        "start", help="Create a resumable intake session."
    )
    start.add_argument(
        "--journey",
        choices=("existing-resume", "scratch"),
        required=True,
    )
    start.add_argument("--out", type=Path, required=True)
    start.add_argument("--source", type=Path)
    start.add_argument("--role", default="")
    start.add_argument("--region", default="")
    start.add_argument("--domain", default="")
    start.add_argument("--job-description", default="")
    start.add_argument(
        "--include-extracted-text",
        action="store_true",
        help="Store full extracted text in the private session file (off by default).",
    )

    status = subparsers.add_parser(
        "status", help="Show intake progress and the next three questions."
    )
    status.add_argument("session", type=Path)
    status.add_argument("--out", type=Path)

    domain = subparsers.add_parser(
        "domain", help="Route a role or job description to question guidance."
    )
    domain.add_argument("query", nargs="+")
    domain.add_argument("--catalog", type=Path)
    domain.add_argument("--out", type=Path)

    validate = subparsers.add_parser("validate", help="Validate a career record.")
    validate.add_argument("record", type=Path)
    validate.add_argument("--out", type=Path)

    approve = subparsers.add_parser(
        "approve", help="Bind final wording after the candidate confirms every fact."
    )
    approve.add_argument("record", type=Path)
    approve.add_argument("--out", type=Path, required=True)

    render = subparsers.add_parser("render", help="Render resume formats.")
    render.add_argument("record", type=Path)
    render.add_argument("--outdir", type=Path, required=True)
    render.add_argument(
        "--formats",
        default="docx,pdf,txt,md",
        help="Comma-separated: docx,pdf,txt,md",
    )
    render.add_argument("--basename")

    extract_cmd = subparsers.add_parser("extract", help="Extract document text.")
    extract_cmd.add_argument("document", type=Path)
    extract_cmd.add_argument("--out", type=Path)

    verify = subparsers.add_parser(
        "verify", help="Compare extracted document text with the career record."
    )
    verify.add_argument("record", type=Path)
    verify.add_argument("document", type=Path)
    verify.add_argument("--out", type=Path)

    interview = subparsers.add_parser(
        "interview", help="Create an evidence-linked interview scaffold."
    )
    interview.add_argument("record", type=Path)
    interview.add_argument("--out", type=Path, required=True)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        if args.command == "start":
            session = new_intake_session(
                journey=args.journey,
                source=args.source,
                role=args.role,
                region=args.region,
                domain=args.domain,
                job_description=args.job_description,
                include_extracted_text=args.include_extracted_text,
            )
            _write_or_print(session, args.out)
            _write_or_print(
                {
                    "session": args.out.name,
                    "journey": session["journey"],
                    "selectedDomain": session["target"]["domain"],
                    "sourceCharacters": sum(
                        item.get("extractedCharacters", 0)
                        for item in session.get("sources", [])
                    ),
                    "nextQuestions": session["nextQuestions"],
                    "privacyNotice": session["privacyNotice"],
                }
            )
            return 0
        if args.command == "status":
            report = intake_status_report(load_record(args.session))
            _write_or_print(report, args.out)
            return 0
        if args.command == "domain":
            report = route_domain(" ".join(args.query), args.catalog)
            _write_or_print(report, args.out)
            return 0
        if args.command == "validate":
            report = validate_record(load_record(args.record))
            _write_or_print(report, args.out)
            return 0 if report["valid"] else 2
        if args.command == "approve":
            approved = approve_record(load_record(args.record))
            _write_or_print(approved, args.out)
            _write_or_print(
                {
                    "approvedRecord": args.out.name,
                    "approvedClaims": len(approved["meta"]["approvedClaimDigests"]),
                    "warning": "Use only after the candidate has reviewed and confirmed every material fact.",
                }
            )
            return 0
        if args.command == "render":
            record = load_record(args.record)
            formats = [item.strip() for item in args.formats.split(",") if item.strip()]
            manifest = render_all(record, args.outdir, formats, args.basename)
            _write_or_print(manifest)
            return 0
        if args.command == "extract":
            _write_or_print(extract_text(args.document), args.out)
            return 0
        if args.command == "verify":
            report = verify_document(load_record(args.record), args.document)
            _write_or_print(report, args.out)
            return 0 if report["pass"] else 3
        if args.command == "interview":
            _write_or_print(interview_markdown(load_record(args.record)), args.out)
            print(args.out.name)
            return 0
    except (
        OSError,
        ValueError,
        RuntimeError,
        TypeError,
        AttributeError,
        UnicodeError,
        json.JSONDecodeError,
    ) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
